"""Reporte de ventas por rango de fechas y exportación a Excel."""

from __future__ import annotations

import datetime as dt
import logging
import tkinter as tk
from tkinter import filedialog, messagebox

from openpyxl import Workbook
from openpyxl.styles import Font

from controller.common import guarded
from model.db_connection import DBConnection
from model.report import SalesReportRow, SalesTotals, rows_from_receipts, totals_from_receipts
from view.sales_report_view import SalesReportView

logger = logging.getLogger(__name__)

EXPORT_HEADERS = (
    "Recibo",
    "Fecha",
    "Hora",
    "Método de pago",
    "Código",
    "Producto",
    "Cantidad",
    "Precio unitario",
    "Subtotal",
)


class SalesReportViewController:
    def __init__(self, parent: tk.Frame, main_controller, db: DBConnection) -> None:
        self.main_controller = main_controller
        self.db = db
        self.rows: list[SalesReportRow] = []
        self.totals: SalesTotals | None = None
        self.view = SalesReportView(parent, self)
        self.view.load_table([])

    def event_back(self) -> None:
        self.main_controller.show_menu()

    def event_dashboard(self) -> None:
        self.main_controller.show_dashboard_view()

    @guarded
    def event_quick_range(self, key: str) -> None:
        today = dt.date.today()
        if key == "today":
            start, end = today, today
        elif key == "yesterday":
            start = end = today - dt.timedelta(days=1)
        elif key == "week":
            start, end = today - dt.timedelta(days=today.weekday()), today
        elif key == "month":
            start, end = today.replace(day=1), today
        else:
            raise ValueError(f"Rango desconocido: {key}")
        self.view.set_dates(start, end)
        self.event_search()

    @guarded
    def event_search(self) -> None:
        start = self.view.get_start_date()
        end = self.view.get_end_date()
        if start > end:
            raise ValueError("La fecha de inicio no puede ser posterior a la fecha final.")
        receipts = self.db.get_receipts_in_range(start, end)
        self.rows = rows_from_receipts(receipts)
        self.totals = totals_from_receipts(receipts)
        self.view.load_table(self.rows)
        self.view.set_totals(self.totals)
        if not receipts:
            messagebox.showinfo("Reporte", "No hay ventas en el rango seleccionado.")

    @guarded
    def event_export(self) -> None:
        if not self.rows:
            messagebox.showinfo("Exportar", "Primero consulte un rango de fechas con ventas.")
            return
        filename = filedialog.asksaveasfilename(
            title="Guardar reporte",
            defaultextension=".xlsx",
            filetypes=[("Excel", "*.xlsx")],
            initialfile=f"reporte_ventas_{dt.datetime.now():%Y%m%d_%H%M%S}.xlsx",
        )
        if not filename:
            return
        workbook = _build_workbook(self.rows, self.totals)
        workbook.save(filename)
        logger.info("Reporte exportado a %s (%s filas)", filename, len(self.rows))
        messagebox.showinfo("Reporte guardado", f"El archivo quedó en:\n{filename}")


def _build_workbook(rows: list[SalesReportRow], totals: SalesTotals | None) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Ventas"
    sheet.append(list(EXPORT_HEADERS))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(
            [
                row.receipt_id,
                row.date,
                row.time.strftime("%H:%M"),
                row.payment_method,
                row.code,
                row.name,
                row.quantity,
                row.unit_price,
                row.total,
            ]
        )
    if totals is not None:
        sheet.append([])
        for label, value in (
            ("Total recaudado", totals.total),
            ("Efectivo", totals.cash),
            ("Tarjeta", totals.card),
            ("Transferencia", totals.transfer),
            ("Recibos", totals.receipt_count),
        ):
            sheet.append([label, "", "", "", "", "", "", "", value])
            sheet.cell(row=sheet.max_row, column=1).font = Font(bold=True)
    sheet.column_dimensions["B"].width = 14
    sheet.column_dimensions["F"].width = 36
    for column in ("H", "I"):
        for cell in sheet[column][1:]:
            cell.number_format = "#,##0"
    return workbook
