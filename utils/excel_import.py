"""Importar y exportar el catálogo de productos en Excel (.xlsx).

El mismo archivo que se exporta sirve de plantilla: tiene los encabezados que
el importador reconoce y una hoja de instrucciones. La importación valida todo
el archivo antes de tocar la base: si hay un solo error, no se importa nada.
"""

from __future__ import annotations

import unicodedata
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from model.product import DEFAULT_MIN_STOCK, TAX_RATES, Product
from utils.formatters import parse_money

HEADERS = ("Código", "Nombre", "Categoría", "Costo", "Precio", "Existencias", "Stock mínimo", "IVA %", "Descripción")
REQUIRED = ("codigo", "nombre", "precio")
MAX_ERRORS = 15
EXAMPLE = Product("7701234567890", "Ejemplo: Gaseosa 350 ml", Decimal(1800), Decimal(2500), 24, "Bebidas", "", True, 6)

INSTRUCTIONS = (
    "Cómo llenar este archivo",
    "",
    "- Una fila por producto. No cambie los encabezados de la primera fila.",
    "- Código: el código de barras o interno; es obligatorio y no puede repetirse.",
    "- Nombre y Precio son obligatorios. El precio es el de venta, con IVA incluido.",
    "- Costo: lo que paga al proveedor. Si no lo sabe, déjelo en 0.",
    "- Existencias: unidades que hay hoy. Para productos que ya existen, déjela vacía si no quiere cambiar el stock.",
    "- Stock mínimo: cuando quedan estas unidades o menos, el sistema avisa (3 si se deja vacío).",
    "- IVA %: 0, 5 o 19.",
    "- Categoría: si no existe, se crea. Vacía = General.",
    "- Los productos con un código que ya existe se actualizan; los demás se crean.",
    "- Borre la fila de ejemplo antes de importar.",
)


@dataclass(frozen=True)
class ImportRow:
    product: Product
    stock: int | None  # None = no cambiar las existencias de un producto que ya existe


@dataclass
class ImportResult:
    rows: list[ImportRow]
    errors: list[str]

    @property
    def ok(self) -> bool:
        return not self.errors


def write_catalog(path: Path, products: list[Product]) -> None:
    """Exporta el catálogo; si está vacío deja una fila de ejemplo para usar el archivo como plantilla."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Productos"
    sheet.append(list(HEADERS))
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F6F95")
    for product in products or [EXAMPLE]:
        sheet.append(
            [
                product.code,
                product.name,
                product.category,
                product.cost,
                product.price,
                product.stock,
                product.min_stock,
                product.tax_rate,
                product.description,
            ]
        )
    widths = (18, 40, 18, 12, 12, 12, 13, 8, 40)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for column in ("D", "E"):
        for cell in sheet[column][1:]:
            cell.number_format = "#,##0"
    sheet.freeze_panes = "A2"

    notes = workbook.create_sheet("Instrucciones")
    for line in INSTRUCTIONS:
        notes.append([line])
    notes["A1"].font = Font(bold=True, size=13)
    notes.column_dimensions["A"].width = 110
    workbook.save(path)


def read_products(path: Path) -> ImportResult:
    """Lee la primera hoja. Devuelve las filas válidas o la lista de errores (nunca ambas)."""
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl lanza varios tipos según el daño del archivo
        return ImportResult([], [f"No se pudo abrir el archivo: {exc}"])
    try:
        sheet = workbook.worksheets[0]
        iterator = sheet.iter_rows(values_only=True)
        header = next(iterator, None)
        if header is None:
            return ImportResult([], ["El archivo está vacío."])
        columns = {_normalize(str(cell)): index for index, cell in enumerate(header) if cell is not None}
        missing = [name for name in REQUIRED if name not in columns]
        if missing:
            return ImportResult(
                [], [f"Faltan las columnas: {', '.join(missing)}. Exporte el Excel para ver el formato."]
            )

        rows: list[ImportRow] = []
        errors: list[str] = []
        seen: dict[str, int] = {}
        for number, values in enumerate(iterator, start=2):
            if values is None or all(value is None or str(value).strip() == "" for value in values):
                continue
            try:
                row = _parse_row(values, columns)
            except ValueError as exc:
                errors.append(f"Fila {number}: {exc}")
                continue
            if row.product.code in seen:
                errors.append(
                    f"Fila {number}: el código {row.product.code} ya está en la fila {seen[row.product.code]}."
                )
                continue
            seen[row.product.code] = number
            rows.append(row)
            if len(errors) >= MAX_ERRORS:
                errors.append("Hay más errores; se muestran los primeros.")
                break
    finally:
        workbook.close()
    if errors:
        return ImportResult([], errors)
    if not rows:
        return ImportResult([], ["El archivo no tiene productos (solo encabezados)."])
    return ImportResult(rows, [])


# ---------------------------------------------------------------------- interno
def _parse_row(values: tuple, columns: dict[str, int]) -> ImportRow:
    def cell(name: str):
        index = columns.get(name)
        if index is None or index >= len(values):
            return None
        value = values[index]
        return None if value is None or str(value).strip() == "" else value

    code = _as_code(cell("codigo"))
    if not code:
        raise ValueError("falta el código.")
    if code == EXAMPLE.code and str(cell("nombre") or "").startswith("Ejemplo"):
        raise ValueError("borre la fila de ejemplo antes de importar.")
    name = str(cell("nombre") or "").strip()
    if not name:
        raise ValueError("falta el nombre.")
    price = _as_decimal(cell("precio"), "precio")
    if price <= 0:
        raise ValueError("el precio debe ser mayor que cero.")
    cost = _as_decimal(cell("costo"), "costo") if cell("costo") is not None else Decimal(0)
    if cost < 0:
        raise ValueError("el costo no puede ser negativo.")
    stock_value = cell("existencias")
    stock = _as_int(stock_value, "existencias") if stock_value is not None else None
    if stock is not None and stock < 0:
        raise ValueError("las existencias no pueden ser negativas.")
    min_value = cell("stock minimo")
    min_stock = _as_int(min_value, "stock mínimo") if min_value is not None else DEFAULT_MIN_STOCK
    if min_stock < 0:
        raise ValueError("el stock mínimo no puede ser negativo.")
    tax_value = cell("iva %")
    tax_rate = _as_decimal(tax_value, "IVA").normalize() if tax_value is not None else Decimal(0)
    if tax_rate not in TAX_RATES:
        raise ValueError("el IVA debe ser 0, 5 o 19.")
    category = str(cell("categoria") or "General").strip() or "General"
    description = str(cell("descripcion") or "").strip()
    product = Product(
        code=code,
        name=name,
        cost=cost,
        price=price,
        stock=stock or 0,
        category=category,
        description=description,
        min_stock=min_stock,
        tax_rate=Decimal(int(tax_rate)),
    )
    return ImportRow(product, stock)


def _normalize(text: str) -> str:
    plain = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    return " ".join(plain.lower().split())


def _as_code(value) -> str:
    """Excel guarda los códigos numéricos como número: 7702004003508.0 debe volver a ser texto."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int) and not isinstance(value, bool):
        return str(value)
    return str(value).strip()


def _as_decimal(value, label: str) -> Decimal:
    if value is None:
        raise ValueError(f"falta el {label}.")
    if isinstance(value, bool):
        raise ValueError(f"el {label} no es un número.")
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))
    try:
        return parse_money(str(value))
    except (ValueError, InvalidOperation) as exc:
        raise ValueError(f"el {label} '{value}' no es un número.") from exc


def _as_int(value, label: str) -> int:
    number = _as_decimal(value, label)
    if number != number.to_integral_value():
        raise ValueError(f"{label} debe ser un número entero.")
    return int(number)
