"""Importación y exportación del catálogo en Excel."""

from __future__ import annotations

from decimal import Decimal

from openpyxl import Workbook, load_workbook

from model.product import Product
from utils import excel_import


def _row(code, name, price=1000, cost=0, stock=1, min_stock=1, tax=0, category="General", description=""):
    """Una fila en el orden de los encabezados exportados."""
    return (code, name, category, cost, price, stock, min_stock, tax, description)


def _write(path, rows, headers=excel_import.HEADERS) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(list(headers))
    for row in rows:
        sheet.append(list(row))
    workbook.save(path)


def test_export_then_import_round_trip(tmp_path):
    products = [
        Product("A1", "Gaseosa", Decimal(2000), Decimal(3000), 5, "Bebidas", "fría", True, 4, Decimal(19)),
        Product("B2", "Pan", Decimal(300), Decimal(500), 0, "General"),
    ]
    path = tmp_path / "inventario.xlsx"
    excel_import.write_catalog(path, products)
    assert load_workbook(path).sheetnames == ["Productos", "Instrucciones"]

    result = excel_import.read_products(path)
    assert result.ok, result.errors
    assert [(r.product.code, r.product.name, r.product.price, r.product.tax_rate, r.stock) for r in result.rows] == [
        ("A1", "Gaseosa", Decimal(3000), Decimal(19), 5),
        ("B2", "Pan", Decimal(500), Decimal(0), 0),
    ]
    assert result.rows[0].product.min_stock == 4 and result.rows[0].product.description == "fría"


def test_empty_catalog_exports_an_example_row_that_must_be_deleted(tmp_path):
    path = tmp_path / "plantilla.xlsx"
    excel_import.write_catalog(path, [])
    result = excel_import.read_products(path)
    assert not result.ok and "ejemplo" in result.errors[0]


def test_numeric_codes_blank_stock_and_colombian_amounts(tmp_path):
    path = tmp_path / "productos.xlsx"
    _write(path, [(7702004003508.0, "Leche", None, "1.800", "2.500", None, None, 5, None)])
    result = excel_import.read_products(path)
    assert result.ok, result.errors
    row = result.rows[0]
    assert row.product.code == "7702004003508" and row.product.category == "General"
    assert row.product.cost == Decimal(1800) and row.product.price == Decimal(2500)
    assert row.stock is None and row.product.stock == 0, "existencias vacías = no cambiar el stock"
    assert row.product.min_stock == 3 and row.product.tax_rate == Decimal(5)


def test_errors_are_reported_per_row_and_nothing_is_imported(tmp_path):
    path = tmp_path / "malo.xlsx"
    _write(
        path,
        [
            _row("A1", "Uno"),
            _row("", "Sin código"),
            _row("A2", ""),
            _row("A3", "IVA raro", tax=7),
            _row("A1", "Repetido"),
            _row("A4", "Precio", price="abc"),
            _row("A5", "Negativo", stock=-1),
        ],
    )
    result = excel_import.read_products(path)
    assert result.rows == []
    assert [error.split(":")[0] for error in result.errors] == [
        "Fila 3",
        "Fila 4",
        "Fila 5",
        "Fila 6",
        "Fila 7",
        "Fila 8",
    ]
    assert "código" in result.errors[0]
    assert "nombre" in result.errors[1]
    assert "IVA" in result.errors[2]
    assert "fila 2" in result.errors[3]
    assert "abc" in result.errors[4]
    assert "negativas" in result.errors[5]


def test_missing_required_columns_is_a_single_clear_error(tmp_path):
    path = tmp_path / "columnas.xlsx"
    _write(path, [("A1", 1000)], headers=("Código", "Precio"))
    result = excel_import.read_products(path)
    assert not result.ok and "nombre" in result.errors[0]


def test_headers_are_matched_without_accents_or_case(tmp_path):
    path = tmp_path / "mayusculas.xlsx"
    headers = (
        "CODIGO",
        "nombre",
        "Categoria",
        "COSTO",
        "precio",
        "Existencias",
        "stock minimo",
        "iva %",
        "descripcion",
    )
    _write(path, [("A1", "Uno", "Aseo", 100, 200, 3, 1, 0, "")], headers=headers)
    result = excel_import.read_products(path)
    assert result.ok and result.rows[0].product.category == "Aseo"


def test_unreadable_file_is_a_clear_error(tmp_path):
    path = tmp_path / "roto.xlsx"
    path.write_bytes(b"esto no es un excel")
    result = excel_import.read_products(path)
    assert not result.ok and "No se pudo abrir" in result.errors[0]
